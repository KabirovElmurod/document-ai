import pdfplumber
from docx import Document
from openpyxl import load_workbook
import pytesseract
from PIL import Image
from app.config import settings


def extract_pdf(file_path: str) -> str:
    """
    PDF fayldan matn ajratish.
    
    Args:
        file_path: PDF fayl yo'li
        
    Returns:
        Ajratilgan matn (string)
    """
    text = ""
    
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise Exception(f"PDF dan matn ajratishda xato: {str(e)}")
    
    return text.strip()


def extract_docx(file_path: str) -> str:
    """
    DOCX fayldan matn ajratish (paragraflar + jadval kataklari).
    
    Args:
        file_path: DOCX fayl yo'li
        
    Returns:
        Ajratilgan matn (string)
    """
    text = ""
    
    try:
        doc = Document(file_path)
        
        # Paragraflar
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        
        # Jadval kataklari
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells])
                if row_text.strip():
                    text += row_text + "\n"
                    
    except Exception as e:
        raise Exception(f"DOCX dan matn ajratishda xato: {str(e)}")
    
    return text.strip()


def extract_xlsx(file_path: str) -> str:
    """
    XLSX fayldan matn ajratish (har varaq qatorlarini matnga aylantirish).
    
    Args:
        file_path: XLSX fayl yo'li
        
    Returns:
        Ajratilgan matn (string)
    """
    text = ""
    
    try:
        wb = load_workbook(file_path, read_only=True)
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text += f"Sheet: {sheet_name}\n"
            
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    text += row_text + "\n"
            
            text += "\n"
            
    except Exception as e:
        raise Exception(f"XLSX dan matn ajratishda xato: {str(e)}")
    
    return text.strip()


def extract_image(file_path: str, lang: str = 'eng+rus') -> str:
    """
    Rasm fayldan matn ajratish (OCR).
    
    Args:
        file_path: Rasm fayl yo'li (png, jpg, jpeg)
        lang: Tesseract til kodi (masalan: 'eng+rus' uchun ingliz va rus)
        
    Returns:
        Ajratilgan matn (string)
    """
    try:
        if settings.TESSERACT_PATH:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_PATH
        
        if settings.TESSDATA_PREFIX:
            pytesseract.pytesseract.tessdata_dir_config = settings.TESSDATA_PREFIX
        
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image, lang=lang)
        return text.strip()
    except Exception as e:
        raise Exception(f"Rasmdan matn ajratishda xato: {str(e)}")


def extract_text(file_path: str, content_type: str) -> str:
    """
    Fayl turi bo'yicha mos extract funksiyasini chaqiradi.
    
    Args:
        file_path: Fayl yo'li
        content_type: Fayl kontent turi (MIME type)
        
    Returns:
        Ajratilgan matn (string)
    """
    if content_type == "application/pdf":
        return extract_pdf(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_xlsx(file_path)
    elif content_type in ["image/png", "image/jpeg", "image/jpg"]:
        return extract_image(file_path)
    else:
        raise Exception(f"Qo'llab-quvvatlanmaydigan fayl turi: {content_type}")

